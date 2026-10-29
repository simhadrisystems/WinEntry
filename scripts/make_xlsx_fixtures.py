#!/usr/bin/env python3
"""Writes the synthetic openpyxl fixture used by XlsxReaderTest (no real business data).

Usage: python3 scripts/make_xlsx_fixtures.py
"""
import datetime
import os

from openpyxl import Workbook

OUT = os.path.join(os.path.dirname(__file__), "..", "app", "src", "test", "resources", "xlsx")

wb = Workbook()
ws = wb.active
ws.title = "Products"
ws.append(["Name", "Code", "Qty", "Price", "Active"])
ws.append(["Sample Item A", "W1249", 12, 450.5, True])
ws.append(["Sample Item B", "B1139L", 0, 120, False])
ws.append(["Telugu name సరుకు", "00123", None, 99.99, True])
ws["C5"] = "=C2+C3"

dates = wb.create_sheet("Dates")
formats = ["yyyy-mm-dd", "dd/mm/yyyy", "d-mmm-yy", "[$-409]d-mmm-yyyy", "dd.mm.yyyy", "mm-dd-yy", "0.00", "General"]
for i, fmt in enumerate(formats, start=1):
    dates.cell(row=i, column=1, value=fmt)
    c = dates.cell(row=i, column=2, value=datetime.datetime(2026, 4, 30) if fmt not in ("0.00", "General") else 46142)
    c.number_format = fmt

wb.save(os.path.join(OUT, "openpyxl_sample.xlsx"))
print("wrote", os.path.join(OUT, "openpyxl_sample.xlsx"))
